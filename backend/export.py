"""Exports for university systems: Moodle XML, Platonus TXT and results as Excel."""
import io
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import scoring


def _text(value) -> str:
    return f'<text>{escape(str(value))}</text>'


def moodle_xml(title: str, questions: list[dict]) -> str:
    """Moodle "Moodle XML format" question bank (multichoice, shortanswer, matching)."""
    out = ['<?xml version="1.0" encoding="UTF-8"?>', "<quiz>"]
    out.append(
        f'  <question type="category"><category>{_text("$course$/TestGen/" + title)}</category></question>'
    )
    for i, q in enumerate(questions, 1):
        name = f"<name>{_text(f'{i}. ' + q['question'][:60])}</name>"
        qtext = f'<questiontext format="plain_text">{_text(q["question"])}</questiontext>'
        if q["type"] == "multiple_choice":
            answers = "".join(
                f'<answer fraction="{100 if opt == q["correct"] else 0}" format="plain_text">{_text(opt)}</answer>'
                for opt in q["options"]
            )
            out.append(
                f'  <question type="multichoice">{name}{qtext}<defaultgrade>1</defaultgrade>'
                f"<single>true</single><shuffleanswers>true</shuffleanswers>"
                f"<answernumbering>abc</answernumbering>{answers}</question>"
            )
        elif q["type"] == "open":
            out.append(
                f'  <question type="shortanswer">{name}{qtext}<defaultgrade>1</defaultgrade>'
                f'<usecase>0</usecase><answer fraction="100" format="plain_text">{_text(q["correct"])}</answer></question>'
            )
        elif q["type"] == "matching":
            subs = "".join(
                f'<subquestion format="plain_text">{_text(term)}<answer>{_text(definition)}</answer></subquestion>'
                for term, definition in q["correct"].items()
            )
            out.append(
                f'  <question type="matching">{name}{qtext}<defaultgrade>1</defaultgrade>'
                f"<shuffleanswers>true</shuffleanswers>{subs}</question>"
            )
    out.append("</quiz>")
    return "\n".join(out)


def platonus_txt(questions: list[dict]) -> str:
    """Platonus import format: the first <variant> is the correct answer.

    Platonus only supports single-choice questions, so other types are skipped.
    """
    blocks = []
    for q in questions:
        if q["type"] != "multiple_choice":
            continue
        variants = [q["correct"]] + [o for o in q["options"] if o != q["correct"]]
        lines = [f"<question>{q['question']}"] + [f"<variant>{v}" for v in variants]
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks) + "\n"


def results_xlsx(title: str, attempts: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Нәтижелер"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    headers = ["№", "Аты-жөні", "Тобы", "Балл", "Барлығы", "%", "Әріптік баға", "GPA", "Дәстүрлі баға",
               "Уақыты (мин)", "Тапсырған күні"]
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B4FCF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, a in enumerate(attempts, 1):
        g = scoring.grade(a["percent"])
        taken = a["taken_at"]
        try:
            taken = datetime.fromisoformat(taken).astimezone().strftime("%d.%m.%Y %H:%M")
        except Exception:
            pass
        minutes = round(a["duration_sec"] / 60, 1) if a.get("duration_sec") else ""
        ws.append([i, a["student_name"], a.get("student_group") or "", a["score"], a["total"],
                   a["percent"], g["letter"], g["gpa"], g["text"], minutes, taken])

    widths = [5, 30, 14, 8, 9, 8, 13, 7, 20, 12, 18]
    for col, w in zip("ABCDEFGHIJK", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
